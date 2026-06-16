import hashlib
import json
import os
import asyncio
from pathlib import Path
from typing import Callable, Optional
from sqlmodel import select
from sqlmodel.ext.asyncio.session import AsyncSession
from models import Document, DocumentFragment
from services.extraction import extract_structured_entities, segment_into_fragments, classify_fragment_basic

DATA_DIR = Path("./data/documents")
DATA_DIR.mkdir(parents=True, exist_ok=True)


def compute_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


async def extract_text(file_path: Path, filename: str) -> str:
    ext = filename.lower().split(".")[-1]
    try:
        if ext == "pdf":
            return await _extract_pdf(file_path)
        elif ext in ("docx", "doc"):
            return await _extract_docx(file_path)
        elif ext in ("txt", "md"):
            return file_path.read_text(encoding="utf-8", errors="replace")
        elif ext in ("xlsx", "xls"):
            return await _extract_xlsx(file_path)
        else:
            try:
                from markitdown import MarkItDown
                md = MarkItDown()
                result = md.convert(str(file_path))
                return result.text_content
            except Exception:
                return file_path.read_text(encoding="utf-8", errors="replace")
    except Exception as e:
        return f"[Error extrayendo texto: {e}]"


async def _extract_pdf(file_path: Path) -> str:
    try:
        from markitdown import MarkItDown
        md = MarkItDown()
        result = md.convert(str(file_path))
        return result.text_content
    except Exception:
        pass
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(str(file_path)) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text_parts.append(t)
        return "\n\n".join(text_parts)
    except Exception:
        pass
    return "[No se pudo extraer texto del PDF]"


async def _extract_docx(file_path: Path) -> str:
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(str(file_path))
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as e:
        return f"[Error con DOCX: {e}]"


async def _extract_xlsx(file_path: Path) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(file_path), read_only=True)
        rows = []
        for sheet in wb.worksheets[:3]:
            for row in list(sheet.iter_rows(values_only=True))[:100]:
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append(" | ".join(cells))
        return "\n".join(rows)
    except Exception as e:
        return f"[Error con XLSX: {e}]"


async def process_document(
    doc_id: int,
    entity_id: int,
    file_path: Path,
    original_name: str,
    doc_type: str,
    session: AsyncSession,
    broadcast: Optional[Callable] = None,
    scoped_domains: Optional[list] = None,
) -> None:
    async def emit(event: str, data: dict = {}):
        if broadcast:
            await broadcast(str(entity_id), {"event": event, "doc_id": doc_id, **data})

    try:
        doc_result = await session.exec(select(Document).where(Document.id == doc_id))
        doc = doc_result.first()
        if not doc:
            return

        doc.status = "processing"
        session.add(doc)
        await session.commit()
        await emit("doc_step", {"step": "Extrayendo texto...", "progress": 10})

        text = await extract_text(file_path, original_name)
        await emit("doc_step", {"step": "Analizando entidades...", "progress": 40})

        entities_found = extract_structured_entities(text)
        doc.extracted_entities = json.dumps(entities_found, ensure_ascii=False)
        await emit("doc_step", {"step": "Segmentando contenido...", "progress": 60})

        raw_fragments = segment_into_fragments(text)
        doc.page_count = len(raw_fragments)
        session.add(doc)
        await session.commit()
        await emit("doc_step", {"step": "Clasificando fragmentos...", "progress": 80})

        for frag_data in raw_fragments:
            classification = classify_fragment_basic(frag_data["content"], doc_type)
            fragment = DocumentFragment(
                document_id=doc_id,
                entity_id=entity_id,
                content=frag_data["content"],
                fragment_type=classification["fragment_type"],
                cobit_hint=classification["cobit_hint"],
                confidence=classification["confidence"],
                page_ref=frag_data.get("page_ref"),
            )
            session.add(fragment)

        await session.commit()
        await emit("doc_step", {"step": "Recalculando cobertura COBIT...", "progress": 90})

        from services.cobit_rules import recalculate_coverage
        coverage_result = await recalculate_coverage(entity_id, session, scoped_domains)

        doc.status = "ready"
        doc.classification_status = "complete"
        session.add(doc)
        await session.commit()

        await emit("doc_complete", {
            "fragments": len(raw_fragments),
            "new_findings": len(coverage_result.get("new_findings", [])),
        })

    except Exception as e:
        doc_result = await session.exec(select(Document).where(Document.id == doc_id))
        doc = doc_result.first()
        if doc:
            doc.status = "error"
            session.add(doc)
            await session.commit()
        await emit("doc_error", {"message": str(e)})
